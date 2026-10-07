"""Tests run on a small made-up dataset in the same layout as the real
financials workbook, so they never depend on licensed data."""

from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook

from ratio_engine import (
    HIGHLIGHT_STYLE,
    IDENTITY_PAIRS,
    INDUSTRY_HIGHLIGHTS,
    INSIGHTS_ENGINE,
    ITEM_COLUMNS,
    RATIO_CATEGORIES,
    UNIVERSAL_RATIOS,
    build_insight_markdown,
    category_scores,
    correlation_matrix,
    cross_industry_medians,
    data_checks,
    describe_correlation,
    format_value,
    highlights_for,
    industry_category_summary,
    industry_panel,
    industry_trend,
    load_financials,
    matrix_styles,
    pair_counts,
    peer_percentile,
    regression_slope,
    style_matrix,
    top_correlations,
)

APP_PATH = str(Path(__file__).resolve().parent.parent / "app.py")
INDUSTRIES = ["Technology / Software", "FMCG"]
YEARS = ["FY2024", "FY2025", "FY2026"]
N_COMPANIES = 12


def make_rows() -> pd.DataFrame:
    rng = np.random.default_rng(0)
    rows = []
    for industry in INDUSTRIES:
        for c in range(N_COMPANIES):
            for y in YEARS:
                rev = rng.uniform(1_000, 10_000)
                cogs = rev * rng.uniform(0.3, 0.7)
                ebitda = rev * rng.uniform(0.1, 0.3)
                ebit = ebitda * 0.8
                equity = rng.uniform(2_000, 8_000)
                debt = equity * rng.uniform(0.0, 1.5)
                interest = debt * 0.08
                ni = (ebit - interest) * 0.75
                assets = equity + debt + rng.uniform(500, 2_000)
                price = rng.uniform(100, 1_000)
                mcap = ni * rng.uniform(10, 40)
                rows.append(
                    {
                        "Industry": industry,
                        "Company": f"{industry.split()[0]} Co {c:02d}",
                        "Capital IQ Ticker": f"NSEI:X{c}",
                        "Period": y,
                        "Period End Date": pd.Timestamp(f"{y[2:]}-03-31"),
                        "Total Revenue": rev,
                        "Cost of Goods Sold": cogs,
                        "Gross Profit": rev - cogs,
                        "EBIT": ebit,
                        "EBITDA": ebitda,
                        "Interest Expense": interest,
                        "Net Income": ni,
                        "Total Current Assets": assets * 0.5,
                        "Total Current Liabilities": assets * rng.uniform(0.2, 0.4),
                        "Cash & ST Investments": assets * 0.1,
                        "Inventory": cogs * rng.uniform(0.05, 0.3),
                        "Total Receivables": rev * rng.uniform(0.1, 0.3),
                        "Total Assets": assets,
                        "Total Debt": debt,
                        "Total Equity": equity,
                        "Cash from Operations": ebitda * 0.7,
                        "Debt Repaid": debt * 0.1,
                        "Dividend per Share": price * rng.uniform(0.0, 0.03),
                        "Share Price (period end)": price,
                        "Market Cap": mcap,
                        "Total Enterprise Value": mcap + debt,
                        "Notes": None,
                    }
                )
    df = pd.DataFrame(rows)
    # Special cases the loader must handle
    first = df.index[(df["Company"] == "Technology Co 00") & (df["Period"] == "FY2026")][0]
    df.loc[first, ["Inventory", "Debt Repaid"]] = 0
    df.loc[first, "Notes"] = "Not reported on Screener, set to 0: inventory, debt repaid"
    loss = df.index[(df["Company"] == "Technology Co 01") & (df["Period"] == "FY2026")][0]
    df.loc[loss, "Net Income"] = -500
    neg_eq = df.index[(df["Company"] == "FMCG Co 00") & (df["Period"] == "FY2025")][0]
    df.loc[neg_eq, "Total Equity"] = -100
    no_int = df.index[(df["Company"] == "FMCG Co 01") & (df["Period"] == "FY2024")][0]
    df.loc[no_int, ["Interest Expense", "Total Debt"]] = 0
    big_div = df.index[(df["Company"] == "FMCG Co 02") & (df["Period"] == "FY2024")][0]
    df.loc[big_div, "Dividend per Share"] = df.loc[big_div, "Share Price (period end)"] * 0.5
    return df


def write_workbook(df: pd.DataFrame, path: Path) -> Path:
    """Same layout as the real file: note row, header row, mnemonic row, data."""
    wb = Workbook()
    wb.active.title = "Instructions"
    ws = wb.create_sheet("Data")
    ws.append(["Test data"])
    ws.append(list(df.columns))
    ws.append([None] * 5 + [f"IQ_{i}" for i in range(len(df.columns) - 5)])
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in row])
    wb.save(path)
    return path


@pytest.fixture(scope="module")
def xlsx_path(tmp_path_factory):
    return write_workbook(make_rows(), tmp_path_factory.mktemp("data") / "financials.xlsx")


@pytest.fixture(scope="module")
def data(xlsx_path):
    return load_financials(xlsx_path)


@pytest.fixture(scope="module")
def pool(data):
    return data.pool


def ratio(data, company, year, name):
    r = data.ratios
    return r.loc[(r["Company"] == company) & (r["Year"] == year), name].iloc[0]


def test_definitions():
    assert len(UNIVERSAL_RATIOS) == 25 and len(set(UNIVERSAL_RATIOS)) == 25
    assert len(RATIO_CATEGORIES) == 5 and all(len(v) == 5 for v in RATIO_CATEGORIES.values())
    assert set(INSIGHTS_ENGINE) == set(INDUSTRY_HIGHLIGHTS)
    for industry, highlights in INDUSTRY_HIGHLIGHTS.items():
        assert len(highlights) == 4 and set(highlights) <= set(UNIVERSAL_RATIOS)
        md = build_insight_markdown(industry)
        assert all(f"**{r}**" in md and INSIGHTS_ENGINE[industry][r] for r in highlights)
    assert all(p <= set(UNIVERSAL_RATIOS) for p in IDENTITY_PAIRS)
    assert len(highlights_for("Unknown Industry")) == 4


def test_loader_shape(data):
    assert data.years == YEARS
    assert list(data.pool) == INDUSTRIES
    assert len(data.ratios) == len(INDUSTRIES) * N_COMPANIES * len(YEARS)
    df = data.pool["FMCG"]["FMCG Co 03"]
    assert list(df.index) == UNIVERSAL_RATIOS
    assert list(df.columns) == ["Ratio Category", *YEARS]


def test_csv_and_xlsx_give_same_ratios(data, tmp_path):
    csv = tmp_path / "financials.csv"
    make_rows().to_csv(csv, index=False)
    from_csv = load_financials(csv)
    pd.testing.assert_frame_equal(
        from_csv.ratios.reset_index(drop=True), data.ratios.reset_index(drop=True), check_exact=False, atol=1e-6
    )


def test_formulas_match_hand_calculation(data):
    raw = data.raw
    row = raw[(raw["Company"] == "FMCG Co 05") & (raw["Year"] == "FY2025")].iloc[0]
    get = lambda name: ratio(data, "FMCG Co 05", "FY2025", name)  # noqa: E731
    assert get("Current Ratio") == pytest.approx(row["Total Current Assets"] / row["Total Current Liabilities"], rel=1e-3)
    assert get("Net Profit Margin") == pytest.approx(100 * row["Net Income"] / row["Total Revenue"], rel=1e-3)
    assert get("Return on Equity (ROE)") == pytest.approx(100 * row["Net Income"] / row["Total Equity"], rel=1e-3)
    assert get("Inventory Turnover") == pytest.approx(row["Cost of Goods Sold"] / row["Inventory"], rel=1e-3)
    assert get("Days Sales Outstanding (DSO)") == pytest.approx(365 * row["Total Receivables"] / row["Total Revenue"], rel=1e-3)
    assert get("Price-to-Earnings (P/E)") == pytest.approx(row["Market Cap"] / row["Net Income"], rel=1e-3)
    assert get("EV/EBITDA") == pytest.approx(row["Total Enterprise Value"] / row["EBITDA"], rel=1e-3)
    dscr = row["EBITDA"] / (row["Interest Expense"] + row["Debt Repaid"])
    assert get("Debt Service Coverage Ratio (DSCR)") == pytest.approx(dscr, rel=1e-3)


def test_placeholders_and_meaningless_ratios_are_blank(data):
    assert any("set to 0" in i for i in data.issues)
    assert pd.isna(ratio(data, "Technology Co 00", "FY2026", "Inventory Turnover"))
    assert pd.isna(ratio(data, "Technology Co 00", "FY2026", "Days Inventory Outstanding (DIO)"))
    assert not pd.isna(ratio(data, "Technology Co 00", "FY2026", "Debt Service Coverage Ratio (DSCR)"))
    assert pd.isna(ratio(data, "Technology Co 01", "FY2026", "Price-to-Earnings (P/E)"))
    assert ratio(data, "Technology Co 01", "FY2026", "Net Profit Margin") < 0
    for name in ["Return on Equity (ROE)", "Debt-to-Equity", "Equity Multiplier", "Price-to-Book (P/B)"]:
        assert pd.isna(ratio(data, "FMCG Co 00", "FY2025", name))
    assert pd.isna(ratio(data, "FMCG Co 01", "FY2024", "Interest Coverage Ratio"))
    assert ratio(data, "FMCG Co 01", "FY2024", "Debt-to-Equity") == 0


def test_data_checks(data):
    checks = data_checks(data)
    assert ((checks["Company"] == "FMCG Co 02") & checks["Issue"].str.startswith("Dividend yield")).any()
    assert ((checks["Company"] == "FMCG Co 00") & checks["Issue"].str.startswith("Negative equity")).any()


def test_matrix_styles_highlight_exactly_four_rows(pool):
    industry = "FMCG"
    df = next(iter(pool[industry].values()))
    styles = matrix_styles(df, INDUSTRY_HIGHLIGHTS[industry])
    lit = [r for r in styles.index if (styles.loc[r] == HIGHLIGHT_STYLE).all()]
    assert sorted(lit) == sorted(INDUSTRY_HIGHLIGHTS[industry])
    assert "Inventory Turnover" in style_matrix(df, INDUSTRY_HIGHLIGHTS[industry], YEARS[0]).to_html()


def test_format_value_units():
    assert format_value("Net Profit Margin", 12.345) == "12.35%"
    assert format_value("Debt-to-Equity", 1.5) == "1.50x"
    assert format_value("Days Sales Outstanding (DSO)", 42.06) == "42.1d"
    assert format_value("Debt-to-Equity", float("nan")) == "—"


def test_peer_percentile():
    peers = pd.Series([1.0, 2.0, 3.0, 4.0, np.nan])
    assert peer_percentile("Current Ratio", 4.0, peers) == 100.0
    assert peer_percentile("Debt-to-Equity", 4.0, peers) == 25.0
    assert pd.isna(peer_percentile("Current Ratio", np.nan, peers))


def test_correlations(pool):
    panel = industry_panel(pool, "Technology / Software")
    assert panel.shape == (N_COMPANIES * len(YEARS), 25)
    assert industry_panel(pool, "FMCG", "FY2026").shape == (N_COMPANIES, 25)
    corr = correlation_matrix(panel, "spearman")
    assert corr.shape == (25, 25)
    # DSO = 365 / receivables turnover, so ranks are perfectly reversed
    assert corr.at["Receivables Turnover", "Days Sales Outstanding (DSO)"] == pytest.approx(-1.0)
    counts = pair_counts(panel)
    assert counts.at["Inventory Turnover", "Inventory Turnover"] == len(panel) - 1
    # Too few paired observations -> not reported
    small = correlation_matrix(industry_panel(pool, "FMCG", "FY2026").head(5))
    assert small.isna().all().all()


def test_top_correlations(pool):
    corr = correlation_matrix(industry_panel(pool, "FMCG"))
    all_pairs = top_correlations(corr, include_identities=True)
    economic = top_correlations(corr, include_identities=False, counts=pair_counts(industry_panel(pool, "FMCG")))
    assert (all_pairs["r"].abs().diff().dropna() <= 1e-12).all()
    assert (economic["Link"] == "Economic").all() and "n" in economic.columns
    assert len(all_pairs) > len(economic)


def test_describe_correlation_wording():
    assert "go **up** too" in describe_correlation("A", "B", 0.8)
    assert "go **down**" in describe_correlation("A", "B", -0.5)
    assert "independently" in describe_correlation("A", "B", 0.05)
    assert "cannot be measured" in describe_correlation("A", "B", float("nan"))


def test_regression_slope():
    x = pd.Series([1.0, 2.0, 3.0, np.nan])
    assert regression_slope(x, 2 * x + 1) == pytest.approx(2.0)
    assert pd.isna(regression_slope(pd.Series([1.0, 1.0, 1.0]), pd.Series([1.0, 2.0, 3.0])))


def test_industry_stats(pool):
    year = YEARS[-1]
    summary = industry_category_summary(pool, "Technology / Software", "Profitability", year)
    assert list(summary.index) == RATIO_CATEGORIES["Profitability"]
    assert (summary["P25"] <= summary["Median"]).all() and (summary["Median"] <= summary["P75"]).all()
    assert summary.at["Return on Equity (ROE)", "Companies"] == N_COMPANIES
    assert list(industry_trend(pool, "FMCG", "Net Profit Margin").index) == YEARS
    assert list(cross_industry_medians(pool, "Leverage", year).columns) == INDUSTRIES
    scores = category_scores(pool, "FMCG", "Profitability", year)
    assert len(scores) == N_COMPANIES and scores.between(0, 100).all() and scores.is_monotonic_decreasing


def test_app_runs_on_workbook(xlsx_path, monkeypatch):
    from streamlit.testing.v1 import AppTest

    monkeypatch.setenv("FINANCIALS_PATH", str(xlsx_path))
    at = AppTest.from_file(APP_PATH, default_timeout=60).run()
    assert not at.exception
    for industry in INDUSTRIES:
        at.selectbox[0].select(industry).run()
        assert not at.exception
        assert len(at.metric) == 4
        assert INSIGHTS_ENGINE[industry]["title"] in at.info[0].value
    at.selectbox[0].select("Technology / Software").run()
    at.selectbox[1].select("Technology Co 01").run()  # company with a loss year
    assert not at.exception
    for category in RATIO_CATEGORIES:
        at.radio[0].set_value(category).run()
        assert not at.exception
    at.radio[1].set_value(at.radio[1].options[1]).run()  # single-year scope
    at.radio[2].set_value("Pearson").run()
    at.toggle[0].set_value(True).run()
    assert not at.exception


def test_app_without_data_asks_for_upload(tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest

    monkeypatch.setenv("FINANCIALS_PATH", str(tmp_path / "missing.xlsx"))
    at = AppTest.from_file(APP_PATH, default_timeout=60).run()
    assert not at.exception
    assert "No data loaded" in at.info[0].value


def test_item_columns_match_template():
    from openpyxl import load_workbook

    template = Path(__file__).resolve().parent.parent / "data" / "capiq_template.xlsx"
    header = [c.value for c in load_workbook(template, read_only=True)["Data"][2]]
    assert set(ITEM_COLUMNS) <= set(header)
